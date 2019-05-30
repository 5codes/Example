'''
Author: C5
Time: 2018.05.01
Contact: 17755105543
Base Function: Read datas from a docx file,and export to a xls files.
v2: add col 'Room' and export position of datas in word file
'''
#coding=utf-8
import docx
import xlrd,xlwt,xlutils   #xlrd and xlutils are not used in this Script
import datetime
#Parameters defined
para_text = []
document = docx.Document('Example_text-bak.docx')   #引号内写入需处理的文件名
para = document.paragraphs

def get_docx():
    #transfer to text
    for i in para:
        para_text.append(i.text)

def get_amount_of_datas():
    #count the amount of these datas
    num = 1
    Count = 0
    CountBlank = 0
    Sequences = []
    SequencesBlank = []
    SequencesBlank1 = []

    for i in para:
        if i.text == "我：":
            Count += 1
            Sequences.append(para.index(i))
        elif i.text == "":
            CountBlank += 1
            SequencesBlank.append(para.index(i))

    for i in SequencesBlank:
        if i > Sequences[-1]:
            SequencesBlank1.append(i)
            break
        elif i == Sequences[num] - 1:
            SequencesBlank1.append(i)
            num += 1
    return Sequences, SequencesBlank1, Count

def get_datas(Sequences,SequencesBlank1):
    #Store these datas in a list and a dictionary
    DATAs = []
    Datas = []
    for sequence in Sequences:
        tmp = Sequences.index(sequence)
        for i in para_text[sequence+1:SequencesBlank1[tmp]]:
            i = i.split('：')
            if len(i) == 2:
                if i[0] == "日期":
                    date = i[1]
#                   print('Date:\t' + date)
                elif i[0] == "医院":
                    Hospitor = i[1]
#                   print('Hospitor:\t' + Hospitor)
                elif i[0] == "科室":
                    Room = i[1]
#                   print('Room:\t' + Room)
                elif i[0] == "处方医生":
                    Doctor = i[1]
#                   print('Doctor:\t' + Doctor)
                elif i[0] == "患者":
                    Patient = i[1]
#                   print('Patient:\t' + Patient)
                elif i[0] == "性别":
                    Sex = i[1]
#                   print('Sex:\t' + Sex)
                elif i[0] == "年龄":
                    Age = i[1]
#                   print('Age:\t' + Age)
                elif i[0] == "癌种":
                    Cancel = i[1]
#                   print('Cancel:\t' + Cancel)
                elif i[0] == "剂量":
                    Amount = i[1]
#                   print('Amount:\t' + Amount)
                elif i[0] == "购买盒数":
                    Number_box = i[1]
#                   print('Number:\t' + Number_box)
                elif i[0] == "购买盒数":
                    Sale_site = i[1]
#                   print('Sale_site:\t' + Sale_site)
                elif i[0] == "业务员":
                    Saler = i[1]
#                   print('Saler:\t' + Saler)
        new_DATA = {'Date': date, 'Hospitor': Hospitor, 'Room': Room,'Doctor': Doctor, 'Patient': Patient, 'Sex': Sex, 'Age': Age,'Cancel': Cancel, 'Amount': Amount, 'Number_box': Number_box, 'Saler': Saler}
        DATAs.append(new_DATA)
        new_data = [date, Hospitor, Room, Doctor, Patient, Sex, Age, Cancel, Amount, Number_box, Saler]
        Datas.append(new_data)
    return DATAs, Datas

def write_to_xls(Datas):
    #write datas to excel
    ExcelName = datetime.datetime.now().strftime("%Y%m%d") + '.xls'
    workbook = xlwt.Workbook(encoding='utf-8')
    booksheet = workbook.add_sheet('Add_datas', cell_overwrite_ok=True)
    for row in range(len(Datas)):
        print("数据行："+ str(Sequences[row]))
        print("输出第"+ str(row+1) + "条数据")
        for col in range(len(Datas[row])):
            booksheet.write(row, col, Datas[row][col])
    workbook.save(ExcelName)
    print("处理完成！！！")

get_docx()
Sequences, SequencesBlank1, Count = get_amount_of_datas()
print('共有数据条数：\t'+str(Count))
print(Sequences)
print(SequencesBlank1)
try:
    DATAs, Datas = get_datas(Sequences,SequencesBlank1)
    write_to_xls(Datas)
except IndexError:
    print("文档格式需处理！！！")

input('Press the enter key to exit.')  #屏显暂停，enter才会退出


